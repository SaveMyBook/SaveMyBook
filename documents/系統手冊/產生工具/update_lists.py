"""重建「圖目錄」「表目錄」。

兩份清單各包成一個 TOC \\c 功能變數（圖目錄 \\c "圖"、表目錄 \\c "表"），圖表說明內須有同名 SEQ 欄位才會被收錄：
原稿前幾章的說明已有可見的 SEQ 編號欄位，其餘說明（圖表號為文字）另加隱藏的 SEQ 欄位（\\h 不顯示結果）。
Word 更新功能變數後即依實際版面重排頁碼。欄位內預先放入的清單只是更新前的顯示內容，頁碼無法在沒有排版引擎的
情況下精算：原有的圖表沿用原頁碼，再加上前面新增圖表估計佔用的頁數；新增的圖表接續前一筆估算。
"""
import copy, re
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from doctools import W, make_run

LIST_STYLE = 'table of figures'
NEW_PAGES = {'圖': 1.0, '表': 0.6}
_norm = lambda s: re.sub(r'\s+', '', s)


def _entries(doc):
    return [p for p in doc.paragraphs if p.style.name == LIST_STYLE]


CAPTION = re.compile(r'^(圖|表) \d+(-\d+)+ \S')


def _is_caption(p):
    """圖表說明多為 Caption 樣式，少數沿用內文樣式（如圖 1-2-7），改以緊鄰的圖片或表格判斷。"""
    text = p.text.strip()
    if not CAPTION.match(text) or p.style.name == LIST_STYLE: return False
    if p.style.name == 'Caption': return True
    prev, nxt = p._p.getprevious(), p._p.getnext()
    if text.startswith('圖'):
        return prev is not None and prev.find('.//' + qn('a:blip')) is not None
    return nxt is not None and nxt.tag == W('tbl')


def _captions(doc):
    return [p for p in doc.paragraphs if _is_caption(p)]


def capture(doc):
    """在任何修改之前呼叫：記下每個原有圖表說明所在的頁碼。"""
    pages = {}
    for p in _entries(doc):
        label, _, page = p.text.partition('\t')
        if page.strip().isdigit(): pages[_norm(label)] = int(page.strip())
    kept = {}
    for p in _captions(doc):
        page = pages.get(_norm(p.text))
        if page is not None: kept[p._p] = page
    return kept


def _entry(template, label, number, title, page):
    p = copy.deepcopy(template)
    for child in list(p):
        if child.tag != W('pPr'): p.remove(child)
    for text in (label, f' {number} ', title):
        p.append(make_run(None, text))
    r = make_run(None, str(page))
    r.insert(0, r.makeelement(W('tab'), {}))
    p.append(r)
    return p


def run(doc, kept):
    rows, offset, last = [], 0.0, 0
    for p in _captions(doc):
        label, number, title = p.text.strip().split(' ', 2)
        if p._p in kept:
            page = kept[p._p] + round(offset)
        else:
            offset += NEW_PAGES[label]
            page = last + (1 if label == '圖' else 0)
        page = max(page, last)
        last = page
        rows.append((label, number, title, page))

    entries = _entries(doc)
    template = copy.deepcopy(entries[0]._p)
    for br in template.iter(W('lastRenderedPageBreak')):
        br.getparent().remove(br)
    counts = {}
    for label in ('圖', '表'):
        block = [e._p for e in entries if e.text.startswith(label)]
        anchor = block[0]
        for lab, number, title, page in rows:
            if lab == label:
                anchor.addprevious(_entry(template, lab, number, title, page))
        for el in block: el.getparent().remove(el)
        counts[label] = sum(1 for r in rows if r[0] == label)
    for label in ('圖', '表'):
        block = [e._p for e in _entries(doc) if e.text.startswith(label)]
        ppr = block[0].find(W('pPr'))
        at = list(block[0]).index(ppr) + 1 if ppr is not None else 0
        for i, r in enumerate((_fld('begin'), _instr(f' TOC \\h \\z \\c "{label}" '), _fld('separate'))):
            block[0].insert(at + i, r)
        block[-1].append(_fld('end'))
    counts['hidden_seq'] = _tag_captions(doc)
    return counts


def _fld(kind):
    r = OxmlElement('w:r'); c = OxmlElement('w:fldChar'); c.set(W('fldCharType'), kind); r.append(c)
    return r


def _instr(text):
    r = OxmlElement('w:r'); t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = text; r.append(t)
    return r


def _seq_labels(p_el):
    instr = ''.join(x.text or '' for x in p_el.iter(W('instrText')))
    return set(re.findall(r'SEQ\s+(\S+)', instr))


def _tag_captions(doc):
    """沒有 SEQ 欄位的圖表說明補上隱藏的 SEQ 欄位，供 TOC \\c 收錄。
    同一個標題 2 小節內，隱藏欄位若排在可見的 SEQ 編號之前會使其編號加一，因此遇到這種情形即中止。"""
    added, hidden_in_section = 0, set()
    for p in doc.paragraphs:
        if p.style.name.lower() == 'heading 2':
            hidden_in_section = set(); continue
        if not _is_caption(p): continue
        label = p.text.strip()[0]
        if label in _seq_labels(p._p):
            assert label not in hidden_in_section, f'visible SEQ after hidden SEQ in same section: {p.text}'
            continue
        p._p.extend([_fld('begin'), _instr(f' SEQ {label} \\h '), _fld('end')])
        hidden_in_section.add(label); added += 1
    return added

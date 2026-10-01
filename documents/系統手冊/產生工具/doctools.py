import copy, re, struct
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Emu

W = lambda tag: qn('w:' + tag)

def text_of(el):
    return ''.join(x.text or '' for x in el.iter(W('t')))

def _run_template(p):
    r = p.find(W('r'))
    return copy.deepcopy(r.find(W('rPr'))) if r is not None and r.find(W('rPr')) is not None else None

def make_run(rpr, text):
    r = OxmlElement('w:r')
    if rpr is not None: r.append(copy.deepcopy(rpr))
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    r.append(t)
    return r

def set_para(p_el, text):
    """保留段落格式與第一個 run 的字型，改成單一 run 的新文字。"""
    rpr = _run_template(p_el)
    for child in list(p_el):
        if child.tag != W('pPr'): p_el.remove(child)
    p_el.append(make_run(rpr, text))

def set_cell(tc, lines):
    if isinstance(lines, str): lines = lines.split('\n')
    ps = tc.findall(W('p'))
    tpl = ps[0]
    ppr = tpl.find(W('pPr')); rpr = _run_template(tpl)
    for p in ps: tc.remove(p)
    for line in lines:
        p = OxmlElement('w:p')
        if ppr is not None: p.append(copy.deepcopy(ppr))
        p.append(make_run(rpr, line))
        tc.append(p)

def tcs(tr): return tr.findall(W('tc'))

def row_by_key(tbl, key):
    for tr in tbl.findall(W('tr')):
        c = tcs(tr)
        if c and text_of(c[0]).strip() == key: return tr
    raise KeyError(key)

def clone_row_after(ref_tr, template_tr=None):
    new = copy.deepcopy(template_tr if template_tr is not None else ref_tr)
    ref_tr.addnext(new)
    return new

def png_size(path):
    with open(path, 'rb') as f:
        head = f.read(24)
    return struct.unpack('>II', head[16:24])

def fit_size(png_path, max_w, max_h, dpi=220):
    """以圖檔原始尺寸（220 dpi）放入，超出版面才等比縮小，避免小圖被放大。"""
    w, h = png_size(png_path)
    cx, cy = int(w / dpi * 914400), int(h / dpi * 914400)
    scale = min(1, max_w / cx, max_h / cy)
    return int(cx * scale), int(cy * scale)

def replace_image(doc, p_el, png_path, max_w, max_h):
    blip = p_el.find('.//' + qn('a:blip'))
    part = doc.part.related_parts[blip.get(qn('r:embed'))]
    part._blob = open(png_path, 'rb').read()
    cx, cy = fit_size(png_path, max_w, max_h)
    for ext in p_el.iter(qn('wp:extent')):
        ext.set('cx', str(cx)); ext.set('cy', str(cy))
    for ext in p_el.iter(qn('a:ext')):
        if ext.get('cx') is not None:
            ext.set('cx', str(cx)); ext.set('cy', str(cy))

def new_picture_para(doc, template_p, png_path, max_w, max_h):
    cx, cy = fit_size(png_path, max_w, max_h)
    para = doc.add_paragraph()
    para.add_run().add_picture(png_path, width=Emu(cx), height=Emu(cy))
    el = para._p
    tpl_ppr = template_p.find(W('pPr'))
    if tpl_ppr is not None:
        old = el.find(W('pPr'))
        if old is not None: el.remove(old)
        el.insert(0, copy.deepcopy(tpl_ppr))
    el.getparent().remove(el)
    return el

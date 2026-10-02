import sys, re, hashlib, zipfile, collections
sys.path.insert(0, '.')
import docx
from docx.oxml.ns import qn
import update_lists, update_part1, update_fk, schema_model

path = sys.argv[1]
d = DOC = docx.Document(path)
body = d.element.body
ok = True
def fail(msg):
    global ok
    ok = False
    print('FAIL', msg)

caps = update_lists._captions(d)
figs = [p for p in caps if p.text.strip().startswith('圖')]
tabs = [p for p in caps if p.text.strip().startswith('表')]
print('captions: 圖', len(figs), '表', len(tabs))

# 1. 每個圖說前一段是圖片；圖檔內容與預期 PNG 相同
expected = {num: png for num, (png, _, _) in update_part1.FIGS.items()}
for n, (png, _, _) in enumerate(update_part1.ER_FIGS, start=1):
    expected[f'8-1-{n}'] = png
checked = 0
for p in figs:
    prev = p._p.getprevious()
    blip = prev.find('.//' + qn('a:blip')) if prev is not None else None
    if blip is None:
        fail(f'no image before {p.text}'); continue
    num = p.text.split(' ')[1]
    if num in expected:
        blob = d.part.related_parts[blip.get(qn('r:embed'))].blob
        if hashlib.sha1(blob).digest() != hashlib.sha1(open(expected[num], 'rb').read()).digest():
            fail(f'image mismatch {num}')
        checked += 1
print('figure images verified against PNG:', checked, '/', len(expected))
# PNG 內嵌的 PlantUML 原始碼須與同名 .puml 相同，否則圖檔未依最新原始碼重繪
def _puml_src(b):
    import zlib
    i = 8
    while i < len(b):
        ln = int.from_bytes(b[i:i + 4], 'big')
        if b[i + 4:i + 8] == b'iTXt':
            key, rest = b[i + 8:i + 8 + ln].split(b'\0', 1)
            if key == b'plantuml':
                rest = rest[2:].split(b'\0', 1)[1].split(b'\0', 1)[1]
                return (zlib.decompress(rest) if b[i + 8 + len(key) + 1] else rest).decode()
        i += 12 + ln
_norm = lambda s: re.sub(r'\s+', '', s.split('@enduml')[0])
stale = [png.rsplit('/', 1)[1] for png in expected.values()
         if _norm(_puml_src(open(png, 'rb').read()) or '') != _norm(open(png[:-4] + '.puml', encoding='utf-8').read())]
if stale: fail(f'PNG not rendered from current .puml: {stale}')
print('PNG rendered from current .puml:', len(expected) - len(stale), '/', len(expected))
# 表說明後一個元素是表格
for p in tabs:
    nxt = p._p.getnext()
    if nxt is not None and nxt.tag != qn('w:tbl') and ''.join(x.text or '' for x in nxt.iter(qn('w:t'))).startswith('註'):
        nxt = nxt.getnext()
    if nxt is None or nxt.tag != qn('w:tbl'): fail(f'no table after {p.text}')

# 2. 編號連續
for label, group in (('圖', figs), ('表', tabs)):
    seq = collections.OrderedDict()
    for p in group:
        num = p.text.split(' ')[1]
        head, _, last = num.rpartition('-')
        seq.setdefault(head, []).append(int(last))
    bad = {k: v for k, v in seq.items() if v != list(range(1, len(v) + 1))}
    if bad: fail(f'{label} numbering gaps {bad}')
    print(label, 'groups', len(seq), 'numbering continuous' if not bad else 'NOT continuous')

# 3. 圖片尺寸
# 上限依版面可用範圍計算（複評版排版會把活動圖放到接近整頁）：寬取版心寬並容許 0.1cm，高扣除圖說約 1cm。
_sect = body.find(qn('w:sectPr'))
_pg, _mar = _sect.find(qn('w:pgSz')), _sect.find(qn('w:pgMar'))
_twip = lambda el, k: int(el.get(qn('w:' + k)))
MAX_W = int((_twip(_pg, 'w') - _twip(_mar, 'left') - _twip(_mar, 'right')) * 635 + 36000)
MAX_H = int((_twip(_pg, 'h') - _twip(_mar, 'top') - _twip(_mar, 'bottom')) * 635 - 360000)
fig_imgs = set()
over, sizes = [], []
for p in figs:
    img = p._p.getprevious()
    fig_imgs.add(img)
    for ext in img.iter(qn('wp:extent')):
        cx, cy = int(ext.get('cx')), int(ext.get('cy'))
        sizes.append((cx, cy, p.text.split(' ')[1]))
        if cx > MAX_W or cy > MAX_H: over.append((p.text.split(' ')[1], round(cx / 360000, 2), round(cy / 360000, 2)))
if over: fail(f'figure images over page area: {over}')
w = max(sizes); h = max(sizes, key=lambda t: t[1])
print('figure images', len(sizes), 'over page area (%.1fx%.1fcm):' % (MAX_W / 360000, MAX_H / 360000), len(over), '| widest %.2fcm (%s) tallest %.2fcm (%s)' % (w[0] / 360000, w[2], h[1] / 360000, h[2]))
other = []
for el in body.iterchildren():
    if el in fig_imgs: continue
    for ext in el.iter(qn('wp:extent')):
        cx, cy = int(ext.get('cx')), int(ext.get('cy'))
        if cx > MAX_W or cy > MAX_H: other.append((round(cx / 360000, 2), round(cy / 360000, 2)))
print('non-figure images over page area (appendix page scans):', len(other), sorted(set(other)))

# 4. 資料表 Meta data
META = re.compile(r'^表 8-2-(\d+) (\S+) ')
metas = [p for p in tabs if META.match(p.text.strip())]
nums, names, cells = [], [], []
for p in metas:
    m = META.match(p.text.strip())
    nums.append(int(m.group(1))); names.append(m.group(2))
    tbl = p._p.getnext()
    tr0 = tbl.find(qn('w:tr'))
    tc = tr0.findall(qn('w:tc'))
    cells.append(int(''.join(x.text or '' for x in tc[3].iter(qn('w:t')))))
    en = ''.join(x.text or '' for x in tc[1].iter(qn('w:t')))
    if en != m.group(2): fail(f'meta name mismatch {m.group(2)} vs {en}')
print('meta tables', len(metas))
if nums != list(range(1, len(metas) + 1)): fail('meta caption numbers not continuous')
if cells != nums: fail('資料表編號 cells not matching captions')
if names != sorted(names): fail('meta tables not alphabetical')
print('meta caption numbers 1..%d continuous: %s; 資料表編號 match: %s; alphabetical: %s' % (
    len(metas), nums == list(range(1, len(metas) + 1)), cells == nums, names == sorted(names)))
for want in ('book_deposits',):
    print(want, 'present' if want in names else 'MISSING', names.index(want) + 1 if want in names else '')
oi = metas[names.index('order_items')]._p.getnext()
print('order_items.pre_deposited', 'present' if 'pre_deposited' in ''.join(x.text or '' for x in oi.iter(qn('w:t'))) else 'MISSING')

# 5. 圖目錄／表目錄與圖表說明一致
entries = [p.text.split('\t')[0] for p in d.paragraphs if p.style.name == 'table of figures']
want = [p.text.strip() for p in figs] + [p.text.strip() for p in tabs]
print('list entries', len(entries), 'match captions:', entries == want)
if entries != want: fail('lists differ from captions')

# 6. 禁用字詞（本文、表格、頁首頁尾、註腳）
bad_words = re.compile(r'migration|未執行|舊版|降級|遷移', re.I)
z = zipfile.ZipFile(path)
hits = []
for name in z.namelist():
    if name.startswith('word/') and name.endswith('.xml') and re.search(r'(document|header\d*|footer\d*|footnotes|endnotes|comments)\.xml$', name):
        xml = z.read(name).decode('utf-8')
        text = ''.join(re.findall(r'<w:t(?: [^>]*)?>([^<]*)</w:t>', xml))
        instr = ''.join(re.findall(r'<w:instrText(?: [^>]*)?>([^<]*)</w:instrText>', xml))
        for m in bad_words.finditer(text + '\n' + instr):
            s = (text + '\n' + instr)
            hits.append((name, s[max(0, m.start() - 20): m.end() + 20]))
print('forbidden word hits', len(hits))
for h in hits: fail(f'forbidden {h}')
# 7. ER 圖（8-1）：圖名、說明與 ER_FIGS 一致
er = [p for p in figs if p.text.strip().startswith('圖 8-1-')]
print('ER figures', len(er), '/', len(update_part1.ER_FIGS))
if len(er) != len(update_part1.ER_FIGS): fail('ER figure count')
for p, (png, title, text) in zip(er, update_part1.ER_FIGS):
    if p.text.strip().split(' ', 2)[2] != title: fail(f'ER title {p.text.strip()} != {title}')
    desc = ''.join(x.text or '' for x in p._p.getprevious().getprevious().iter(qn('w:t')))
    if desc != text: fail(f'ER description mismatch {p.text.strip()}')
print('ER titles:', ' / '.join(p.text.strip().split(' ', 1)[1] for p in er))

# 8. 外鍵欄：外鍵說明所列欄位標示 V，其餘留空；外鍵說明涵蓋 schema.prisma 所有關聯
schema = schema_model.load(schema_model.FINAL_API)
fk_total = v_total = 0
for p in metas:
    name = META.match(p.text.strip()).group(2)
    cols, fk_head, fks = update_fk.parse(p._p.getnext())
    listed = {fk for _, fk in fks}
    fk_cols = {c for c, _, _ in listed}
    fk_total += len(fks)
    if fks and fk_head is None: fail(f'{name}: 外鍵說明 rows without heading')
    if set(schema[name]['fks']) != listed:
        fail(f'{name}: 外鍵說明 {sorted(listed)} != schema {sorted(schema[name]["fks"])}')
    for c in fk_cols - set(cols): fail(f'{name}: FK column {c} has no row')
    for c, tr in cols.items():
        mark = ''.join(x.text or '' for x in tr.findall(qn('w:tc'))[-1].iter(qn('w:t'))).strip()
        if mark not in ('', 'V'): fail(f'{name}.{c}: unexpected 外鍵 value {mark!r}')
        if c in fk_cols and mark != 'V': fail(f'{name}.{c}: FK without V')
        if c not in fk_cols and mark: fail(f'{name}.{c}: V on non-FK column')
        v_total += mark == 'V'
print('外鍵說明 rows', fk_total, '| schema relations', sum(len(m['fks']) for m in schema.values()),
      '| columns marked V', v_total, '| tables with FKs', sum(1 for m in schema.values() if m['fks']))
# 9. 資料字典涵蓋最終版 schema 的每一張表與每一個欄位
listed_tables = {META.match(p.text.strip()).group(2): p._p.getnext() for p in metas}
missing_tables = sorted(set(schema) - set(listed_tables))
if missing_tables: fail(f'schema tables without meta table: {missing_tables}')
col_diff = []
for name, tbl in listed_tables.items():
    if name not in schema: fail(f'meta table not in schema: {name}'); continue
    cols, _, _ = update_fk.parse(tbl)
    want = [c for c, _ in schema[name]['cols']]
    if set(want) != set(cols): col_diff.append((name, sorted(set(want) - set(cols)), sorted(set(cols) - set(want))))
for d in col_diff: fail(f'columns differ {d}')
print('schema tables', len(schema), '| meta tables', len(listed_tables), '| missing', missing_tables, '| column mismatches', len(col_diff))
# 10. 圖目錄／表目錄為 TOC \c 功能變數，每個圖表說明都有同名 SEQ 欄位可供收錄
list_ps = [p for p in DOC.paragraphs if p.style.name == 'table of figures']
for label in ('圖', '表'):
    block = [p._p for p in list_ps if p.text.startswith(label)]
    instr = ''.join(x.text or '' for x in block[0].iter(qn('w:instrText')))
    kinds = [x.get(qn('w:fldCharType')) for p in block for x in p.iter(qn('w:fldChar'))]
    if f'TOC \\h \\z \\c "{label}" \\f ' not in instr or kinds != ['begin', 'separate', 'end']:
        fail(f'{label}目錄 is not a TOC \\c field: {instr!r} {kinds}')
no_seq = [p.text.strip()[:12] for p in figs + tabs if p.text.strip()[0] not in update_lists._seq_labels(p._p)]
if no_seq: fail(f'captions without SEQ field: {no_seq[:5]}')
print('lists as TOC \\c fields; captions without SEQ:', len(no_seq))
print('ALL OK' if ok else 'SOME CHECKS FAILED')

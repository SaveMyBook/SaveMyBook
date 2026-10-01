import copy
from doctools import *
from doclib import captioned_tables
from docx.oxml import OxmlElement

def set_row(tr, values):
    for tc, v in zip(tcs(tr), values):
        if v is not None: set_cell(tc, v)

def add_row(tbl, values, after=None):
    rows = tbl.findall(W('tr'))
    ref = after if after is not None else rows[-1]
    new = clone_row_after(ref, rows[-1] if after is None else ref)
    set_row(new, values)
    return new

def replace_in_cell(tr, col, old, new):
    tc = tcs(tr)[col]
    lines = [text_of(p) for p in tc.findall(W('p'))]
    joined = '\n'.join(lines)
    assert old in joined, (old, joined[:80])
    set_cell(tc, joined.replace(old, new).split('\n'))

def manual_table(doc, heading):
    body = list(doc.element.body.iterchildren())
    for i, el in enumerate(body):
        if el.tag == W('p') and text_of(el).strip() == heading:
            for nxt in body[i + 1:]:
                if nxt.tag == W('tbl'): return nxt
    raise KeyError(heading)

def step_set(tbl, step, text): set_cell(tcs(row_by_key(tbl, str(step)))[1], text)

def steps_replace(tbl, texts):
    rows = tbl.findall(W('tr'))
    tpl = rows[1]
    for r in rows[1:]: tbl.remove(r)
    for i, text in enumerate(texts, start=1):
        tr = copy.deepcopy(tpl)
        set_row(tr, [str(i), text])
        tbl.append(tr)

def step_append(tbl, text):
    rows = tbl.findall(W('tr'))
    n = int(text_of(tcs(rows[-1])[0]).strip()) + 1
    add_row(tbl, [str(n), text])

def run(doc):
    T = captioned_tables(doc)
    t = lambda prefix: next(v._tbl for k, v in T.items() if k.startswith(prefix + ' '))

    # 系統架構敘述
    for p in doc.paragraphs:
        if p.text.startswith('行動應用客戶端：'):
            set_para(p._p, '行動應用客戶端：基於Flutter框架建構之跨平台應用程式，負責介面渲染與使用者操作邏輯，並透過RESTful API與後端進行非同步資料交換；聊天功能另以Socket.IO（WebSocket）接收新訊息、已讀與輸入中之即時推送。')
        elif p.text.startswith('RESTful API服務：'):
            set_para(p._p, 'RESTful API服務：基於Node.js執行環境與Express.js框架建構。負責處理來自代理層之HTTP請求，執行包含JWT身分驗證、權限管控、書籍交易匹配、設備狀態同步等核心業務邏輯；同一行程另以Socket.IO提供聊天即時推送，並於背景排程執行訂單自動完成、語意索引更新與書籍資料補齊。另以獨立之裝置API（/api/device/v1）服務智慧書櫃終端，並以Device憑證與使用者權杖區隔驗證。此層級與外部網路邏輯隔離，確保核心商業邏輯免於直接暴露。')

    # App 中賣家的書籍列表頁標題為「書籍管理」
    for node in doc.element.body.iter(W('t')):
        if node.text and '「我的商品」' in node.text:
            node.text = node.text.replace('「我的商品」', '「書籍管理」')

    # 資料庫結構以 prisma/schema.prisma 同步建立
    for p in doc.paragraphs:
        for r in p.runs:
            if '進行資料庫建模與遷移' in r.text:
                r.text = r.text.replace('進行資料庫建模與遷移', '進行資料庫建模與結構同步')

    # 3-2-5 技術平台
    x = t('表 3-2-5')
    ai = row_by_key(x, '後端邏輯層／（伺服器端）') if False else None
    for tr in x.findall(W('tr')):
        c = [text_of(tc) for tc in tcs(tr)]
        if len(c) > 1 and c[1] == 'AI服務':
            set_row(tr, [None, None, 'OpenAI／Google Gemini／DeepSeek（可切換供應商）', '提供上架內容審核、AI客服、AI問書機器人、個人化推薦、書籍資料補齊與交易爭議分析，並以向量嵌入支援語意檢索。'])
        if len(c) > 1 and c[1] == 'API協定':
            set_row(tr, [None, None, 'RESTful API＋Socket.IO（WebSocket）', '提供標準化的介面供App與硬體端進行資料交換；聊天訊息、已讀與輸入中狀態以Socket.IO即時推送。'])

    x = t('表 3-3-1')
    replace_in_cell(row_by_key(x, '核心框架與套件'), 1, 'Express.js、', 'Express.js、Socket.IO、')

    # 9-1-1 元件清單
    x = t('表 9-1-1')
    rows = x.findall(W('tr'))
    def row_starting(prefix):
        return next(r for r in rows if text_of(tcs(r)[0]).startswith(prefix))
    replace_in_cell(row_starting('Express路由層'), 2, '24組', '29組')
    r = row_starting('服務層 services/orders/')
    set_cell(tcs(r)[2], '購物車結帳與直接購買、訂單狀態機轉換（含取書確認、訂單自動完成與逾期自動取消）、金流撥款／退款之核心交易邏輯，所有轉態皆於Prisma交易內以讀取時狀態為條件式更新。')
    r = row_starting('服務層 services/chat/（')
    set_cell(tcs(r)[0], '服務層 services/chat/（messages, transfers, rooms, groups, controls, risk）')
    set_cell(tcs(r)[2], '聊天訊息、群組管理、封鎖/靜音控制、站內轉帳/請款，以及聊天防詐評分、傳送確認與管理員警示之邏輯。')
    r = row_starting('服務層 services/ai/（')
    set_cell(tcs(r)[0], '服務層 services/ai/（moderation, listing-assist, support, book-chat, recommend, enrich, dispute-assist, catalog-search, knowledge, runner）')
    set_cell(tcs(r)[2], 'AI供應商轉接、上架內容審核、AI輔助上架、AI客服、AI問書機器人、依依據分組之個人化推薦、書籍資料補齊、交易爭議分析、每日用量與費用控管。')
    r = row_starting('jobs/scheduler.js')
    set_cell(tcs(r)[2], '排程作業：裝置清理與聊天通知清理(24h)、帳號刪除清算(1h)、資料庫備份(依需求)、預約到期(5min)、訂單自動完成與逾期取消(10min)、書櫃存書逾期暫停與提醒取回(10min)、書櫃作業逾時處理與提交續做(15s)、書櫃裝置連線檢查(60s)、書櫃挑戰碼清理(10min)與過期事件及配對請求清理(24h)、OAuth暫存清理(10min)、語意索引更新(10min)、書籍資料補齊(30min)、上傳檔案清掃(24h)。')
    r = row_starting('features/selling, books（Flutter）')
    set_cell(tcs(r)[2], '上架/編輯/銷售管理（含掃描書櫃存書與取回之入口）/AI輔助上架/ISBN掃描/書籍瀏覽收藏畫面。')
    r = row_starting('features/chat（Flutter）')
    set_cell(tcs(r)[2], '聊天列表、聊天室（含預約協商、站內轉帳、AI問書機器人、防詐提醒與傳送確認）、群組管理畫面。')
    r = row_starting('lib/prisma.js')
    set_cell(tcs(r)[0], 'lib/prisma.js, lib/upload.js, lib/ai.js, lib/maintenance.js')
    set_cell(tcs(r)[2], 'Prisma Client單例、Multer上傳設定、AI供應商HTTP轉接、維護模式中介。')
    r = row_starting('services/chat/aliases.js')
    set_cell(tcs(r)[0], 'services/chat/aliases.js, chat/history.js, chat/typing.js, chat/codec.js')
    set_cell(tcs(r)[2], '聊天暱稱、歷史可見範圍判定、正在輸入狀態、系統訊息編解碼。')
    last = rows[-1]
    for values in [
        ('services/realtime.js', '後端', 'Socket.IO即時推送：以登入權杖驗證連線並依會員分房，於HTTP回應送出（交易已提交）後推送聊天室變動與輸入中狀態，每5分鐘重新驗證連線。'),
        ('services/chat/risk.js', '後端', '聊天防詐規則評分（內容正規化、跨訊息偵測、帳號因素加權）、傳送確認、風險紀錄與管理員防詐警示。'),
        ('services/ai/lexical.js, ai/semantic.js, ai/catalog-search.js, ai/knowledge.js', '後端', '關鍵字檢索與向量語意檢索，合併排序，供AI客服、AI選書、個人化推薦與相似書籍使用；語意索引保存於ai_embeddings。'),
        ('services/ai/enrich.js', '後端', '依ISBN書目或AI網路搜尋補齊書籍簡介、作者、出版社與出版日期，只寫入仍空白之欄位並記錄來源。'),
        ('services/ai/dispute-assist.js', '後端', '交易爭議AI分析：比對上架資料、照片與爭議說明，產生摘要、發現與建議（僅供管理員參考）。'),
        ('services/realtime_service.dart（Flutter）', '前端', 'Socket.IO用戶端：登入與回到前景時連線、進入背景時斷線，接收聊天室變動與輸入中事件，並以登入權杖重新驗證。'),
        ('features/chat/chat_risk.dart（Flutter）', '前端', '防詐提醒、頂部高風險橫幅、防詐須知與傳送確認視窗。'),
        ('services/book-deposits.js', '後端', '先行存書之登記與取回（由書櫃作業關門後提交，或經管理員確認之手動回報）、存書期間之書櫃變更與重新上架限制、存書滿7天暫停販售、每3天提醒取回、滿14天通知管理員，以及後台存書列表與登記取出。'),
        ('features/selling/book_deposit_actions.dart（Flutter）', '前端', '賣家「掃描書櫃存書」「掃描書櫃取回」之入口分流（掃碼、故障備援之手動回報或暫停服務提示），以及購買已存放於書櫃之書籍前之確認視窗。'),
        ('features/admin/admin_cabinet_deposit_screen.dart（Flutter）', '前端', '後台存書列表：依書櫃與逾期篩選，顯示書名、賣家、書櫃、櫃門、存放天數與暫停販售、已通知管理員等標記，可遠端開啟所在櫃門並登記人員取出。'),
    ]:
        last = add_row(x, list(values), after=last)

    # 9-2-1／9-2-2 套件
    add_row(t('表 9-2-1'), ['socket.io', '^4.8.4', 'WebSocket即時推送（聊天訊息、已讀、輸入中）'])
    add_row(t('表 9-2-1'), ['socket.io-client（開發用）', '^4.8', '即時推送之整合測試'])
    add_row(t('表 9-2-2'), ['socket_io_client', '^3.1.6', '聊天即時推送（Socket.IO用戶端）'])

    # 9-2-3 限流
    x = t('表 9-2-3')
    r = row_by_key(x, '通行密鑰註冊選項與註冊')
    set_row(r, ['通行密鑰註冊選項', None, None, None])
    add_row(x, ['通行密鑰重新命名', '15分鐘', '30次', '使用者'], after=r)
    r = row_by_key(x, '身分驗證（含通行密鑰驗證選項）')
    set_row(r, ['身分驗證（各驗證方式分別計算）', None, None, '使用者＋驗證方式'])
    add_row(x, ['通行密鑰驗證選項', '15分鐘', '30次', '使用者'], after=r)
    add_row(t('表 9-2-3'), ['後台交易爭議AI分析', '10分鐘', '30次', '使用者'])
    add_row(t('表 9-2-3'), ['即時連線回報輸入中（Socket.IO）', '1秒', '1次', '連線'])

    # 9-2-4 錯誤代碼
    x = t('表 9-2-4')
    set_row(row_by_key(x, 'DISPUTE_WINDOW_PASSED'), [None, None, '訂單已完成，或已逾取書後24小時之爭議申請期限'])
    anchor = row_by_key(x, 'BOOK_HAS_ORDERS')
    for values in [('BOOK_HELD', '409', '書籍預約保留期間，賣家不可編輯或下架'),
                   ('BOOK_LOCKED', '409', '書籍交易中或已售出，賣家不可編輯'),
                   ('CABINET_MAINTENANCE', '400/409', '書櫃維修中或裝置故障，暫停存書與書櫃存取；結帳時為書籍存放之書櫃維修中，暫時無法購買'),
                   ('ORDER_NOT_CANCELLABLE', '400', '賣家已存書（含下單時書籍已在書櫃之訂單），雙方不可自行取消，須申請爭議'),
                   ('BOOK_DEPOSITED', '409', '書籍已登記存放於書櫃，不可重複登記、變更書櫃或由管理員刪除'),
                   ('RETRIEVAL_REQUIRED', '409', '書籍仍存放於書櫃，須先至書櫃掃碼取回才能重新上架'),
                   ('NOT_DEPOSITED', '409', '手動回報取回時書籍未登記存放於書櫃'),
                   ('DEPOSIT_NOT_ALLOWED', '409', '書籍未上架、未公開或已有進行中之訂單，無法先行存書'),
                   ('CABINET_REQUIRED', '400', '書籍尚未指定存放之書櫃，無法先行存書'),
                   ('CABINET_UNAVAILABLE', '400/409', '書櫃已停用；書籍須改選其他書櫃，已配對裝置之書櫃暫停書櫃存取'),
                   ('BOOK_SOLD_IN_CABINET', '409', '書籍已售出並轉入進行中之訂單，不可取回'),
                   ('CABINET_HAS_DEPOSITS', '409', '書櫃仍有先行存放之書籍，須登記取出後才能停用')]:
        anchor = add_row(x, list(values), after=anchor)
    anchor = row_by_key(x, 'CHAT_CONTROLS_UNAVAILABLE')
    for values in [('RISK_CONFIRM_REQUIRED', '409', '訊息含聯絡方式、付款資訊或平台外交易，須經傳送者確認後帶confirm_risk重送'),
                   ('ALERT_HANDLED', '409', '聊天防詐警示已被處理')]:
        anchor = add_row(x, list(values), after=anchor)
    set_row(row_by_key(x, 'AUTH_SOCIAL_UNAVAILABLE'), [None, None, '伺服器未設定社群登入所需之Firebase參數'])
    for code in ('SECURITY_UNAVAILABLE', 'SESSION_REQUIRED', 'CHAT_V2_UNAVAILABLE',
                 'CHAT_V3_UNAVAILABLE', 'CHAT_CONTROLS_UNAVAILABLE'):
        tr = row_by_key(x, code); tr.getparent().remove(tr)
    set_row(row_by_key(x, 'PASSKEY_UNAVAILABLE'), [None, None, '伺服器未設定通行密鑰所需之網域參數'])
    add_row(x, ['PASSKEY_ORIGIN_NOT_ALLOWED', '400', '請求來源不在允許之通行密鑰來源（如Android簽署金鑰指紋未設定）'], after=row_by_key(x, 'PASSKEY_VERIFICATION_FAILED'))
    add_row(x, ['BOOKS_IN_CABINET', '400', '尚有書籍存放於書櫃，須先至書櫃掃碼取回才能申請刪除帳號'], after=row_by_key(x, 'OPEN_ORDERS'))

    # 9-2-7 權限鍵
    set_row(row_by_key(t('表 9-2-7'), 'cabinets'), [None, None, '智慧書櫃據點與櫃門維護、裝置配對與撤銷、遠端開啟櫃門、待確認作業與手動回報之處理、存書列表與登記取出'])

    # 9-2-8 路由掛載
    add_row(t('表 9-2-8'), ['/socket.io', 'services/realtime.js', 'Socket.IO即時推送（WebSocket，需登入權杖）'])

    # 10-1-1 後端測試
    x = t('表 10-1-1')
    groups = [
        ('test/ai', 'AI設定、用量與預算通知、同意與隱私揭露、AI客服（混合檢索、知識庫與政策常數一致性、語系）、AI書籍顧問與選書檢索、個人化推薦與上架審核、上架輔助、書籍資料補齊與交易爭議AI分析、書目比對、語意檢索與固定答案評測、結構化輸出與伺服器端查證、注入防護、逾時備援與防止重複、決策紀錄與回饋品質、三家AI服務商轉接器', '495'),
        ('test/auth', None, '72'),
        ('test/chat', '聊天室、訊息與編輯收回、靜音封鎖、群組與管理員、歷史可見範圍、@提及、聊天室轉帳、推播內容、聊天防詐評分與傳送確認、防詐警示、Socket.IO即時推送', '112'),
        ('test/commerce', '書籍上下架與違規鎖定、編輯鎖定、ISBN查詢與限流判斷、AI上架審核與審核事件、管理員書籍管理、先行存書、存書逾期暫停與取回、取回待確認期間暫停販售、後台存書列表、書櫃維修狀態、訂單與結算、直接購買、取書確認與自動完成、逾期自動取消、預約保留、購物車與錢包、客服工單與圖片附件、檢舉與交易爭議（以訂單編號申請）', '302'),
        ('test/link-preview', None, '39'),
        ('test/passkeys', '通行密鑰註冊/登入/身分驗證、挑戰值單次使用與逾時、來源與RP ID檢查（含Android簽署金鑰）、簽章計數倒退防重放、帳號列舉防護、應用程式關聯檔案', '64'),
        ('test/platform', '狀態中介層、帳號安全、帳號隱私與刪除、客服附件之匯出與匿名化、通知中心（排除聊天訊息）與推播、後台權限與會員管理、會員等級管理、操作紀錄、備份、加密編號、輸入驗證、資料庫關聯健檢', '181'),
    ]
    for key, desc, n in groups:
        set_row(row_by_key(x, key), [None, desc, n])
    add_row(x, ['test/uploads', '圖片檔案遺失時的處理、失效圖片欄位清理', '13'])
    add_row(x, ['合計', '共10組測試群組，全數通過', '1,536'])

    # 10-2-1 代表性測試
    x = t('表 10-2-1')
    set_row(row_by_key(x, '訂單申訴逾24小時不可受理'), ['訂單爭議逾24小時不可受理', '訂單已取書且picked_up_at已超過24小時，或訂單已完成，申請爭議', None, None])
    set_row(row_by_key(x, '同一使用者對同訂單不可重複申訴'), ['同一使用者對同訂單不可重複申請爭議', '已有pending爭議案件時再次申請', None, None])
    for values in [
        ('存書後不可自行取消訂單', '訂單已存書時買家呼叫取消', '回傳400 ORDER_NOT_CANCELLABLE', '通過'),
        ('取書滿24小時自動完成並撥款', '訂單picked_up_at超過24小時後執行排程', '訂單轉completed並撥款給賣家', '通過'),
        ('預約保留期間賣家不可編輯', '書籍有未到期之confirmed預約時賣家編輯', '回傳409 BOOK_HELD', '通過'),
        ('含聯絡方式之訊息須確認後才送出', '傳送「加我line abc123」', '回傳409 RISK_CONFIRM_REQUIRED；帶confirm_risk後回傳201', '通過'),
        ('24小時內3則高風險訊息建立防詐警示', '同一會員連續傳送索取驗證碼之訊息', '建立警示並通知管理員，第4則累加觸發次數', '通過'),
        ('被拒絕之請求不推送即時事件', '傳送需確認之訊息而被409拒絕', '對方不會收到chat:room事件', '通過'),
        ('書目來源被限流時不記錄查無資料', 'Google Books回429且Open Library查無此書', '改由AI網路搜尋，搜尋也失敗時不寫入補齊紀錄，待下次排程重試', '通過'),
        ('AI未確認ISBN相符時不採用簡介', '模型回覆matched為false或未附來源網址', '記錄為查無資料，不寫入簡介', '通過'),
        ('書籍已存於書櫃時訂單直接成立為已存書', '書籍登記存書後，買家以同一書櫃結帳', '訂單狀態為deposited，買家可立即取書且無法自行取消', '通過'),
        ('部分書籍未存書時訂單維持待存書', '同一訂單僅部分書籍已登記存書時結帳', '訂單為pending_deposit，通知賣家存入其餘書籍', '通過'),
        ('存書滿7天暫停販售，取回後恢復上架', '存書超過7天仍無訂單時執行排程，再由賣家至書櫃掃碼取回', '書籍改為removed並通知賣家；取回後恢復on_sale', '通過'),
        ('存書期間不可變更書櫃', '已登記存書之書籍修改cabinet_id', '回傳409 BOOK_DEPOSITED', '通過'),
        ('書籍已轉入訂單時不可取回', '存書之書籍被下單後，賣家送出取回', '回傳409 BOOK_SOLD_IN_CABINET', '通過'),
        ('書櫃仍有存書時不可停用', '管理員停用仍有先行存書之書櫃', '回傳409 CABINET_HAS_DEPOSITS', '通過'),
    ]:
        add_row(x, list(values))

    # 10-2-2 App 測試檔
    x = t('表 10-2-2')
    rows = x.findall(W('tr'))
    header, tpl = rows[0], rows[1]
    existing = {text_of(tcs(r)[0]).strip(): text_of(tcs(r)[3]) for r in rows[1:]}
    for r in rows[1:]: x.remove(r)
    files = [
        ('layout_overflow_test.dart', 2, 2010, '143個畫面×14種語系/寬度組合之版面溢位檢查，另含會員中心與錢包之2倍字級檢查'),
        ('book_deposit_test.dart', 50, 53, '登記存書與回報取回、暫停販售標示、存書期間鎖定書櫃、購買已在書櫃之書籍、後台存書列表與登記取出、存書通知導向、按鈕文字不截斷'),
        ('social_login_test.dart', 41, 48, None),
        ('ai_test.dart', 45, 45, 'AI設定序列化與表單、用量圖表與預算通知、上架輔助回應與資料來源標示、上架審核資料、AI同意與重新同意、AI客服、推薦依據分組解析'),
        ('cabinet_flow_test.dart', 41, 41, '書櫃掃碼流程：定位權限、掃碼、確認項目、輸入比對數字、開門後完成或取消、結果顯示、接續進行中之作業與書櫃連結，以及各入口依存取模式顯示掃碼或手動回報'),
        ('cabinet_messages_test.dart', 25, 25, '書櫃錯誤代碼與作業結果之訊息、進入掃碼前之狀態檢查、定位必要條件與精確位置要求'),
        ('passkey_platform_test.dart', 25, 25, '通行密鑰平台層參數補齊、系統錯誤分類、可用性偵測、刪除後通知系統'),
        ('cabinet_api_test.dart', 22, 22, '書櫃作業API（建立作業、比對、結束與取消）、故障備援之手動回報、後台書櫃裝置API（配對、遠端開櫃、櫃門操作、作業與事件列表）'),
        ('chat_v3_test.dart', 21, 21, None),
        ('passkey_flow_test.dart', 20, 20, '後台驗證預設通行密鑰、新增與管理通行密鑰、登入頁入口與流程中停用其他入口'),
        ('cabinet_session_test.dart', 19, 19, '書櫃作業資料模型、比對數字格式、訂單與書籍之書櫃欄位、後台資料模型、倒數計算與輪詢、即時推播過濾'),
        ('admin_cabinet_device_test.dart', 17, 17, '後台書櫃裝置管理：輸入配對碼配對、撤銷裝置、待確認櫃門處理、待確認作業與手動回報、遠端開櫃與數字比對、通知導向'),
        ('chat_link_preview_test.dart', 17, 17, None),
        ('crop_geometry_test.dart', 14, 17, None),
        ('my_reservations_test.dart', 8, 17, '我的預約：已保留與待賣家回覆之排序與剩餘時間、取消預約、沿用直接購買流程購買、開啟與賣家之聊天室、空白訊息、會員中心快捷列之已保留筆數、書籍管理之篩選順序與審核中，以及各語系兩個分頁之版面'),
        ('ai_book_chat_test.dart', 11, 11, 'AI書籍顧問對話、書卡與追問建議、上架輔助新增欄位、未同意AI資料處理時先顯示同意畫面'),
        ('ai_resilience_test.dart', 11, 11, 'AI訊息識別碼與逾時重送、背景產生推薦之重新讀取、不感興趣之復原'),
        ('profile_notifications_test.dart', 11, 11, '會員中心入口、通知分類（不含聊天）、全部已讀與清除依分類執行'),
        ('ai_listing_cost_test.dart', 9, 9, '上架輔助兩階段之權杖與定價、依書況切換AI售價、共用書目快取清除、推薦理由顯示'),
        ('image_fallback_test.dart', 9, 9, None),
        ('admin_level_test.dart', 8, 8, '會員等級門檻規則、階梯預覽與拖曳排序'),
        ('ai_decisions_test.dart', 8, 8, 'AI處理紀錄資料模型、重新開啟對話時還原推薦理由與轉接提示、追問建議、後台處理結果統計'),
        ('ai_feedback_loop_test.dart', 8, 8, 'AI回覆評價（有幫助／沒有幫助）、品質報表、審核卡照片與駁回類別、推薦書磚點擊'),
        ('book_detail_actions_test.dart', 6, 8, '直接購買、編輯鎖定、相似書籍與AI整理簡介之標示'),
        ('cabinet_code_test.dart', 7, 7, '書櫃QR Code解析、系統相機開啟之書櫃連結僅開啟掃描頁'),
        ('dispute_window_test.dart', 6, 6, '取書後24小時爭議申請期限、不可申請爭議之狀態'),
        ('ticket_attachments_test.dart', 6, 6, '客服工單附件上傳與顯示'),
        ('chat_risk_test.dart', 5, 5, '防詐提醒排序與去重、輸入列高度不隨輸入跳動'),
        ('order_flow_ui_test.dart', 5, 5, '訂單紀錄：購買訂單與銷售訂單主頁籤之待處理數量與狀態篩選、訂單卡片內容、取書後完成訂單與申請爭議、待存書之完成存書與取消訂單；書籍管理之已被預約、已售出與已完成書籍不可編輯或取消上架'),
        ('admin_moderation_content_test.dart', 4, 4, '後台內容審核：爭議案件之佐證照片、被檢舉訊息之內容、傳送者與前後文，收回或刪除之訊息顯示中性提示'),
        ('ai_image_prep_test.dart', 4, 4, None),
        ('ai_review_labels_test.dart', 4, 4, '上架審核類別、AI用量功能名稱、錯誤代碼與連線測試項目之中文標籤'),
        ('book_eviction_test.dart', 4, 4, '書籍被刪除後清除本機收藏、購物車與最近瀏覽'),
        ('cabinet_partial_deposit_test.dart', 4, 4, '部分存書：確認項目時提示本次僅能存入部分書籍，完成畫面顯示已存入本數與其餘待存入書籍'),
        ('dispute_ai_panel_test.dart', 4, 4, '爭議AI分析面板：載入保存之分析、建議與可信度、依據照片、內容遭阻擋時停用分析'),
        ('image_crop_view_test.dart', 4, 4, None),
        ('admin_verification_test.dart', 3, 3, None),
        ('takedown_display_test.dart', 3, 3, '強制下架之顯示：後台停止公開顯示之提示與列表、賣家書籍依狀態顯示違規說明'),
        ('ai_retrieval_status_test.dart', 2, 2, '語意檢索狀態顯示（模型、涵蓋率、同步與暫停時間）'),
        ('description_line_breaks_test.dart', 2, 2, '書籍簡介中字面之換行符號還原為實際換行'),
        ('dispute_screen_test.dart', 2, 2, '爭議申請頁：由訂單進入時帶入並鎖定訂單編號、手動輸入訂單編號之格式'),
        ('edit_book_publish_date_test.dart', 2, 2, '編輯書籍時保留僅精確至年或月之出版日期'),
        ('group_nickname_test.dart', 2, 2, '群組暱稱'),
        ('in_app_banner_test.dart', 2, 2, 'App內通知橫幅排隊與同對象合併'),
        ('password_characters_test.dart', 2, 2, None),
        ('sell_condition_carry_test.dart', 2, 2, 'AI判斷之書況帶入上架表單'),
        ('sms_code_expiry_test.dart', 2, 2, '簡訊驗證碼逾時'),
        ('payment_pin_setup_test.dart', 1, 1, None),
        ('verification_flow_test.dart', 1, 1, None),
    ]
    for name, declared, executed, desc in files:
        tr = copy.deepcopy(tpl)
        set_row(tr, [name, str(declared), f'{executed:,}', desc if desc else existing.get(name, '')])
        x.append(tr)

    for p in doc.paragraphs:
        if p.text.startswith('前端（Flutter）另有'):
            set_para(p._p, '後端10組測試群組共1,536項測試案例全數通過。前端（Flutter）另有49個測試檔、531個test／testWidgets宣告，實際展開執行2,563項測試並全數通過，涵蓋密碼字元檢查、圖片裁切幾何運算、爭議申請期限與爭議申請頁、Passkey流程、社群登入、連結預覽、AI功能、訂單紀錄與我的預約、書櫃掃碼流程與比對數字輸入、書櫃存書（含部分存書）與取回、後台書櫃裝置管理與內容審核、防詐提醒、版面溢位等關鍵邏輯與畫面，詳細清單如下表。')

    # 11-1-1 系統元件
    x = t('表 11-1-1')
    set_row(row_by_key(x, '應用伺服器'), [None, 'Node.js 20（LTS）＋ Express 5 ＋ Socket.IO 4', '承載savemybook_api，處理所有RESTful API請求與聊天即時推送；以PM2單一行程執行'])
    set_row(row_by_key(x, '反向代理／邊緣網路'), [None, None, 'TLS憑證終止、CDN快取、WAF防護與流量轉發；/socket.io/路徑須升級為WebSocket，/kiosk（模擬書櫃網頁）須轉發至API'])
    set_row(row_by_key(x, 'AI語言模型服務'), [None, 'OpenAI／Google Gemini／DeepSeek', '上架審核、AI客服、AI問書機器人、個人化推薦、書籍資料補齊、交易爭議分析與語意檢索之向量嵌入'])

    # 11-2-1 後端安裝
    x = t('表 11-2-1')
    set_row(row_by_key(x, '4'), [None, '建立資料庫：於MariaDB建立空白資料庫，執行 npx prisma db push 依 prisma/schema.prisma 建立全部資料表'])
    set_row(row_by_key(x, '5'), [None, '匯入初始資料：執行 prisma/seed.sql，建立服務條款、隱私權政策與常見問題（可重複執行，不會覆蓋既有資料）'])
    replace_in_cell(row_by_key(x, '8'), 1, '與資料庫結構版本', '與資料庫結構是否與prisma/schema.prisma一致')
    step7 = row_by_key(x, '7')
    for tr in x.findall(W('tr'))[1:]:
        n = text_of(tcs(tr)[0]).strip()
        if n.isdigit() and int(n) >= 8: set_cell(tcs(tr)[0], str(int(n) + 1))
    add_row(x, ['8', '設定NGINX即時推送：新增location /socket.io/，proxy_pass至API（http://127.0.0.1:3000），並設定proxy_http_version 1.1、Upgrade與Connection標頭及proxy_read_timeout 75s；未設定時聊天改以輪詢運作'], after=step7)

    # 11-2-3 環境變數
    set_row(row_by_key(t('表 11-2-3'), 'PASSKEY_ORIGINS'), [None, None, 'https://savemybook.today', None])
    set_row(row_by_key(t('表 11-2-3'), 'GOOGLE_BOOKS_API_KEY'), [None, '建議', None, 'ISBN查詢配額；未設定時共用額度常被限流（429），書籍資料補齊將改由AI網路搜尋'])

    # 移除 Migration 相關表格（標題、表名、表格）
    body = list(doc.element.body.iterchildren())
    for i, el in enumerate(body):
        if el.tag == W('p') and text_of(el).strip() in ('Migration 執行順序', 'Migration 內容詳表（002～016）'):
            cap, tbl = body[i + 1], body[i + 2]
            assert text_of(cap).startswith('表 11-2-') and tbl.tag == W('tbl')
            for e in (el, cap, tbl): e.getparent().remove(e)

    # 使用手冊
    m = manual_table(doc, '瀏覽與搜尋書籍')
    step_set(m, 3, '首頁「為您推薦」依推薦依據分組顯示（如「與已收藏的《…》相關」「『考試用書』類別推薦」「更多推薦」）；「最近瀏覽」呈現近期查看過之商品。')
    step_set(m, 4, '點擊書籍卡片進入商品詳情頁，價格下方為書籍資訊卡（作者、出版社、出版日期、ISBN、上架日期），其下為可展開之內容簡介、書櫃據點與「相似書籍」；可收藏、分享、聯絡賣家、加入購物車或直接購買。')
    m = manual_table(doc, '上架書籍（賣家）')
    step_append(m, '上架後若簡介、作者、出版社或出版日期未填，系統會依ISBN於背景自動補齊；由AI依書目整理之簡介，於書籍頁以灰色小字標示「由AI依書目整理」。')
    step_append(m, '上架後可至所選之智慧書櫃點擊「掃描書櫃存書」先行存書，買家下單後即可直接取書，操作方式詳見「智慧書櫃存書與取書」。')
    m = manual_table(doc, '編輯與下架商品')
    step_set(m, 2, '於書籍詳情頁點擊「編輯書籍」可修改售價、書況、照片等資訊；交易中、已售出或預約保留期間之商品無法編輯或下架，書籍存放於書櫃期間無法變更書櫃。')
    step_append(m, '書籍存放於書櫃期間下架後，須先至書櫃點擊「掃描書櫃取回」取回書籍，才能重新上架。')
    m = manual_table(doc, '聊天協商、預約與轉帳')
    step_set(m, 2, '買家可於聊天室內對該書點擊「預約書籍」發起保留請求，賣家同意後保留期間其他人無法購買，賣家也無法編輯或下架；保留到期或取消後會通知收藏此書的使用者。')
    step_append(m, '訊息即時送達並顯示對方輸入中狀態。對方訊息含站外聯絡方式、私下付款或常見詐騙話術時，訊息下方會顯示防詐提醒，高風險時聊天室頂部顯示警示橫幅，可查看「防詐須知」或檢舉；自己傳送聯絡方式或付款資訊時，系統會先請您確認。')
    m = manual_table(doc, '加入購物車與結帳')
    step_set(m, 2, '勾選欲購買之商品後點擊「結帳」，系統將要求輸入交易密碼或完成生物辨識驗證，並以錢包餘額即時扣款；款項由平台保管，訂單完成後才撥給賣家。')
    step_append(m, '只買一本時，可於商品詳情頁點擊「直接購買」，確認金額並完成驗證後即成立訂單，不需加入購物車。')
    step_set(m, 3, '結帳完成後，系統會通知對應賣家於七天內將書籍存入指定智慧書櫃；若訂單內之書籍皆已存放於書櫃，訂單直接轉為「已存書」，可立即取書。')
    m = manual_table(doc, '智慧書櫃存書與取書')
    steps_replace(m, [
        '請於書櫃營業時間內前往書櫃，並允許App存取位置資訊；書櫃螢幕閒置時顯示QR Code，約每30秒更換一次，每組僅能使用一次。',
        '【賣家】訂單成立後7天內，於訂單詳情或「銷售紀錄」點擊「掃描書櫃存書」；書籍上架後亦可於「書籍管理」、「銷售紀錄」的「販售中」分頁或書籍頁點擊「掃描書櫃存書」先行存書。【買家】訂單轉為「已存書」後，於訂單詳情或「購買紀錄」點擊「掃描書櫃取書」，或直接開啟「取書」頁掃描。',
        '以App對準書櫃螢幕上之QR Code掃描，App列出您在此書櫃可辦理之項目；確認勾選後，請於60秒內點擊「開啟櫃門」。',
        '書櫃螢幕顯示兩位數字，請於60秒內在手機輸入。數字僅有一次輸入機會，輸入錯誤或逾時時本次作業取消，須重新掃描；請勿將數字告知他人。',
        '櫃門開啟後，依App與書櫃螢幕指示之櫃門（如A02）放入或取出書籍，關上櫃門後於手機點擊「完成」；倒數時間為30秒，每多一扇櫃門增加15秒（最長90秒），倒數結束時系統亦自動完成。如欲放棄本次作業，請於手機點擊「取消」，訂單與存書狀態不會變更。',
        '櫃門關閉後系統才登記存書或取書，並通知交易對方；書櫃螢幕顯示「作業完成」後回到待掃描畫面。',
        '【賣家】先行存書滿7天仍無人購買時，書籍暫停販售並通知您取回，之後每3天提醒一次，滿14天仍未取回將通知書櫃管理員。請至書櫃點擊「掃描書櫃取回」，依相同步驟取回書籍；因逾期而暫停販售之書籍取回後自動恢復上架。書籍已售出、等待買家取書時不可取回。',
        '【買家】商品詳情頁顯示「書籍已存放於書櫃，下單後即可取書」時，結帳前App會提示訂單成立後即可取書且無法取消訂單。取書後確認書況無誤點擊「完成訂單」，款項即撥入賣家錢包；取書滿24小時未申請爭議時，系統自動完成訂單。',
        '賣家於訂單成立後7天內未存書，或買家於存書後7天內未取書，系統自動取消訂單並全額退款。',
        '書櫃連線中斷或故障時，App會提示改為手動回報；請依客服指示放入或取出書籍後點擊「改為手動回報」，回報經客服確認後才生效，確認前訂單或書籍顯示「待客服確認」。',
    ])
    m = manual_table(doc, '訂單申訴')
    head = next(p for p in doc.paragraphs if p.text.strip() == '訂單申訴')._p
    set_para(head, '訂單爭議')
    step_set(m, 1, '取書後如發現商品與描述不符，可於取書後24小時內、訂單完成前至訂單詳情頁點擊「申請爭議」，爭議申請頁會自動帶入訂單編號；填寫爭議說明並上傳佐證照片（最多6張）後送出。亦可於錢包頁點擊「爭議處理」，自行輸入訂單編號申請。')
    step_set(m, 2, '送出後訂單轉為「爭議處理中」，暫停自動完成與撥款，由管理員審核（可參考AI分析）後裁決退款、駁回或協調處理，結果將以推播通知雙方。')
    m = manual_table(doc, '通知偏好設定')
    step_append(m, '通知中心依交易、帳號、客服、優惠分類列出通知；聊天訊息僅以推播提醒，未讀數顯示於聊天列表。')
    m = manual_table(doc, '後台管理入口導覽')
    step_set(m, 2, '後台依權限顯示對應功能區塊：會員、內容、訂單、內容審核（檢舉、上架審核、防詐警示）、硬體與營運（書櫃監控、存書列表）、等級、客服、公告、系統維運等。')
    step_append(m, '於「硬體與營運」之「存書列表」，可依書櫃或「逾期」篩選賣家於訂單成立前存放之書籍，查看賣家、書櫃、櫃門、存放天數與暫停販售、已通知管理員等標記；可直接遠端開啟書籍所在之櫃門，人員取出書籍後點擊「登記取出」並確認，書籍隨即下架並通知賣家。存書滿14天之通知可直接開啟逾期清單。')
    m = manual_table(doc, '帳號資料匯出與刪除')
    step_append(m, '尚有書籍存放於書櫃時無法申請刪除帳號，須先至書櫃點擊「掃描書櫃取回」取回書籍。')

    # 智慧書櫃（掃碼存取）
    import cabinet_part4
    cabinet_part4.run(doc, t, manual_table, steps_replace)

    # 9-2 各表依出現順序編號（API路由掛載對照表排在權限鍵定義表之前）
    caps92 = [p for p in doc.paragraphs if p.style.name == 'Caption' and p.text.strip().startswith('表 9-2-')]
    for i, p in enumerate(caps92, start=1):
        rest = p.text.strip().split(' ', 2)[2]
        set_para(p._p, f'表 9-2-{i} {rest}')

    # 圖 1-2-7 原稿誤用內文樣式，改為與其他圖說相同之 Caption 樣式
    caps = [p for p in doc.paragraphs if p.text.strip().startswith('圖 ') and p.style.name != 'table of figures']
    ref_ppr = next(p._p.find(W('pPr')) for p in caps if p.style.name == 'Caption')
    for p in caps:
        if p.style.name == '內容' and p._p.getprevious() is not None and p._p.getprevious().find('.//' + qn('a:blip')) is not None:
            old = p._p.find(W('pPr'))
            if old is not None: p._p.remove(old)
            p._p.insert(0, copy.deepcopy(ref_ppr))

    # 開檔時提示更新目錄與圖表目錄
    settings = doc.settings.element
    if settings.find(W('updateFields')) is None:
        uf = OxmlElement('w:updateFields'); uf.set(W('val'), 'true'); settings.append(uf)
